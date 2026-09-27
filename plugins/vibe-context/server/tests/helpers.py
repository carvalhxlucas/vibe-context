import io

import docx


def make_pdf(pages: list[str]) -> bytes:
    """Build a minimal text PDF, one line of ASCII text per page."""
    count = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(count))
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {count} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, text in enumerate(pages):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>".encode()
        )
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")

    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def make_docx() -> bytes:
    document = docx.Document()
    document.add_heading("Spec", level=1)
    document.add_paragraph("The billing service charges customers monthly.")
    document.add_heading("Scope", level=2)
    document.add_paragraph("Refunds are out of scope for the first release.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Plan"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Pro"
    table.cell(1, 1).text = "20 USD"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
